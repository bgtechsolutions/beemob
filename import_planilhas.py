"""
BEEMOB — Importação das planilhas para o Supabase
==================================================

Lê as três planilhas do cliente e carrega no Supabase:

  Comissionamento.xlsx
    aba 'Corretores'                  -> corretores
    aba 'Cadastros'                   -> proprietarios, inquilinos, imoveis, contratos
    aba 'Pgto Comissões Gerais (2)'   -> comissoes (competência 09/2026)

  Planilha Financeira - *.xlsm  (uma por imóvel)
    aba 'Cadastros Gerais'            -> proprietarios, imoveis, inquilinos, contratos (dados completos)
    abas 'Rec_MMAAAA'                 -> lancamentos
    aba  'Tx Condom.'                 -> condominio_itens
    aba  'Cálculo Rescisão'           -> rescisoes

Os dados detalhados dos .xlsm sobrescrevem os do consolidado, porque o
consolidado só tem nome e valores, enquanto os .xlsm têm CPF, RG,
endereço e dados bancários.

USO
---
O RLS do projeto exige autenticação, então a escrita precisa da
service_role key. Ela NUNCA deve ser commitada nem colada em chat.
Pegue em: Supabase > Project Settings > API > service_role.

    # PowerShell
    $env:SUPABASE_SERVICE_KEY = "cole-aqui"
    python import_planilhas.py              # dry-run: só mostra o que faria
    python import_planilhas.py --commit     # grava de verdade

    # bash
    export SUPABASE_SERVICE_KEY="cole-aqui"
    python import_planilhas.py --commit

PRÉ-REQUISITO: rodar supabase_planilhas.sql no SQL Editor antes.
"""

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

# ---------------------------------------------------------------- config

DOWNLOADS = Path.home() / "Downloads"
XLSX_COMISSAO = DOWNLOADS / "Comissionamento.xlsx"
XLSM_GLOB = "Planilha Financeira - *.xlsm"

ABA_COMISSAO = "Pgto Comissões Gerais (2)"  # regra confirmada pelo cliente

# Rateio da aba 'Pgto Comissões Gerais (2)' — cada coluna soma 100%.
RATEIO = {
    #  função        1º aluguel  recorrente
    "Corretor":     (0.40,       0.10),
    "Captador":     (0.20,       0.10),
    "Gestor":       (0.20,       0.25),
    "Gestor ADM":   (0.00,       0.00),
    "Imobiliaria":  (0.20,       0.55),   # sem acento: casa com o CHECK do banco
}

# O status na planilha vem como 'Ativo' / 'Autônomo'. O banco guarda em
# ASCII minúsculo para não depender de codificação.
STATUS_CORRETOR = {"ativo": "ativo", "autônomo": "autonomo",
                   "autonomo": "autonomo", "inativo": "inativo"}

TAXA_ADM_PADRAO = 0.10  # base do comissionamento recorrente


def env_supabase():
    """Lê URL do .env e a service key do ambiente (nunca do disco)."""
    env_path = Path(__file__).parent / ".env"
    url = None
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            m = re.match(r"VITE_SUPABASE_URL=(.*)", line.strip())
            if m:
                url = m.group(1).strip()
    url = os.environ.get("SUPABASE_URL", url)
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    return url, key


# ---------------------------------------------------------------- helpers

def s(v):
    """
    Texto limpo, ou None.

    Normaliza para NFC porque o Excel devolve acentos decompostos (NFD:
    'o' + circunflexo combinante). O Postgres compara bytes, então
    'Autônomo' em NFD não casa com 'Autônomo' em NFC de um CHECK
    constraint — foi exatamente o que derrubou a carga de corretores.
    """
    if v is None:
        return None
    t = unicodedata.normalize("NFC", str(v)).strip()
    return t or None


def brl(v):
    """Formata em Real: 1234.5 -> '1.234,50'."""
    return f"{v:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def num(v, default=0.0):
    """Número, tolerante a texto e a fórmula não calculada."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace("R$", "").replace(".", "").replace(",", ".")
        t = re.sub(r"[^\d.\-]", "", t)
        try:
            return float(t)
        except ValueError:
            return default
    return default


def inteiro(v):
    n = num(v, None)
    return int(n) if n is not None else None


def dt(v):
    """Data ISO, ou None."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def dias_no_mes(d: date) -> int:
    nxt = date(d.year + (d.month == 12), (d.month % 12) + 1, 1)
    return (nxt - date(d.year, d.month, 1)).days


# ---------------------------------------------------------------- supabase

class Supa:
    def __init__(self, url, key, commit):
        self.url, self.key, self.commit = url, key, commit
        self.contagem = {}

    def upsert(self, tabela, linhas, on_conflict="id"):
        linhas = [r for r in linhas if r]
        if not linhas:
            return

        chaves = [c.strip() for c in on_conflict.split(",")]

        # (1) Deduplica pela chave de conflito. O Postgres recusa o lote
        #     inteiro com "ON CONFLICT DO UPDATE command cannot affect row
        #     a second time" se a mesma chave aparecer duas vezes. Ao juntar,
        #     o registro mais completo prevalece campo a campo.
        vistos = {}
        for r in linhas:
            k = tuple(r.get(c) for c in chaves)
            if k in vistos:
                vistos[k].update({kk: vv for kk, vv in r.items() if vv is not None})
            else:
                vistos[k] = dict(r)
        duplicados = len(linhas) - len(vistos)
        if duplicados:
            print(f"   · {tabela}: {duplicados} duplicata(s) de "
                  f"{on_conflict} consolidada(s)")
        linhas = list(vistos.values())

        self.contagem[tabela] = self.contagem.get(tabela, 0) + len(linhas)
        if not self.commit:
            print(f"   [dry-run] {tabela}: {len(linhas)} registros")
            return

        # (2) O PostgREST exige que todo objeto do lote tenha exatamente as
        #     mesmas chaves ("All object keys must match"). Em vez de
        #     preencher com null — o que sobrescreveria os DEFAULT da tabela —
        #     envia um lote por formato de registro.
        grupos = {}
        for r in linhas:
            grupos.setdefault(tuple(sorted(r)), []).append(r)
        for grupo in grupos.values():
            self._enviar(tabela, grupo, on_conflict)

    def _enviar(self, tabela, linhas, on_conflict):
        endpoint = f"{self.url}/rest/v1/{tabela}?on_conflict={on_conflict}"
        req = urllib.request.Request(
            endpoint,
            data=json.dumps(linhas, ensure_ascii=False).encode("utf-8"),
            headers={
                "Content-Type": "application/json",
                "apikey": self.key,
                "Authorization": f"Bearer {self.key}",
                "Prefer": "resolution=merge-duplicates,return=minimal",
            },
            method="POST",
        )
        try:
            urllib.request.urlopen(req)
            print(f"   ✓ {tabela}: {len(linhas)} registros")
        except urllib.error.HTTPError as e:
            print(f"   ✗ {tabela}: HTTP {e.code} — {e.read().decode()[:400]}")


# ---------------------------------------------------------------- parsers

def ler_corretores(wb):
    """Aba 'Corretores': ID | Nome | CRECI | RG | CPF | Endereço | Telefone | Status"""
    ws = wb["Corretores"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or not s(r[1]):
            continue
        out.append({
            "id": s(r[0]),
            "nome": s(r[1]),
            "creci": s(r[2]),
            "rg": s(r[3]),
            "cpf": s(r[4]),
            "endereco": s(r[5]),
            "telefone": s(r[6]),
            "status": STATUS_CORRETOR.get((s(r[7]) or "ativo").lower(), "ativo"),
        })
    return out


def ler_cadastros(wb):
    """
    Aba 'Cadastros' do consolidado.
    A | CÓDIGO | B NOME/PROPRIETÁRIO | C NOME/INQUILINO | D $ Aluguel
    E $ Condomínio | F $ IPTU | G ENERGIA SOLAR | H $ Total
    I CORRETOR | J CAPTADOR | K Valor Pgto Proprietário
    L Vr Condom. | M Vr. IPTU | N n.º PIX
    """
    ws = wb["Cadastros"]
    proprietarios, inquilinos, imoveis, contratos = [], [], [], []

    # Blocos da aba 'Cadastros':
    #   4..31  imóveis com contrato ativo
    #   33     cabeçalho do bloco de caução
    #   34..47 cauções
    #   52     cabeçalho de 'IMÓVEIS PARA ALUGAR'
    #   53..70 imóveis sem contrato
    for r in ws.iter_rows(min_row=4, max_row=31, values_only=True):
        cod = s(r[0])
        if not cod:
            continue
        # Aceita '2025-002' e também '2026-042.1' — o sufixo .N é como a
        # planilha representa vários imóveis do mesmo proprietário
        # (ex.: ESPÓLIO DE ROMEU CARLOS MOURÃO, com 3 imóveis).
        m = re.match(r"^(\d{4}-\d{3})(?:\.(\d+))?$", cod)
        if not m:
            continue
        cod_prop, seq = m.group(1), m.group(2) or "1"
        imovel_id = f"{cod_prop}{seq}"   # mesma convenção da planilha: =$A$3&1
        contrato_id = cod                # preserva o .N para não colidir

        proprietarios.append({
            "id": cod_prop,
            "nome": s(r[1]) or cod_prop,
            "pix": s(r[13]),
        })

        if s(r[2]):
            inquilinos.append({"id": imovel_id, "nome": s(r[2])})

        imoveis.append({
            "id": imovel_id,
            "proprietario_id": cod_prop,
            "status": "alugado",
            "valor_aluguel": num(r[3]),
            "valor_condominio": num(r[4]),
            "valor_iptu": num(r[5]),
            "energia_solar": num(r[6]),
            "corretor_nome": s(r[8]),
            "captador_nome": s(r[9]),
            "percentual_taxa_adm": TAXA_ADM_PADRAO,
            "valor_adm": round(num(r[3]) * TAXA_ADM_PADRAO, 2),
            "imobiliaria_nome": "BEEMOB",
        })

        contratos.append({
            "id": contrato_id,
            "proprietario_id": cod_prop,
            "inquilino_id": imovel_id if s(r[2]) else None,
            "imovel_id": imovel_id,
            "status": "ativo",
            "percentual_taxa": TAXA_ADM_PADRAO,
            "corretor_nome": s(r[8]),
            "captador_nome": s(r[9]),
        })

    return proprietarios, inquilinos, imoveis, contratos


def ler_caucoes(wb):
    """
    Bloco 'VALOR REFERENTE À CAUÇÃO EM CONTA APLICAÇÃO BEEMOB'
    (aba 'Cadastros', cabeçalho na 33, dados 34..47).
    A código | B inquilino | C valor caução | D adm 5,7% | I complemento | J total
    É dinheiro de terceiros sob custódia — não é receita.
    """
    ws = wb["Cadastros"]
    out = []
    for r in ws.iter_rows(min_row=34, max_row=47, values_only=True):
        cod = s(r[0])
        if not cod or not re.match(r"^\d{4}-\d{3}", cod):
            continue
        valor = num(r[2])
        if valor <= 0:
            continue
        m = re.match(r"^(\d{4}-\d{3})(?:\.(\d+))?$", cod)
        cod_prop, seq = m.group(1), m.group(2) or "1"
        out.append({
            "contrato_id": cod,
            "imovel_id": f"{cod_prop}{seq}",
            "inquilino_nome": s(r[1]),
            "valor_caucao": valor,
            "percentual_adm": 0.057,
            "valor_adm": num(r[3]),
            "valor_complementar": num(r[8]),
            "valor_total": num(r[9]),
            "devolvido": False,
        })
    return out


def ler_imoveis_disponiveis(wb):
    """
    Bloco 'IMÓVEIS PARA ALUGAR' (aba 'Cadastros', linhas 51-70):
    imóveis da carteira sem contrato ativo. 'VAGO' = sem proprietário
    identificado na planilha.
    """
    ws = wb["Cadastros"]
    proprietarios, imoveis = [], []
    for r in ws.iter_rows(min_row=53, max_row=70, values_only=True):
        cod = s(r[0])
        if not cod or not re.match(r"^\d{4}-\d{3}$", cod):
            continue
        nome = s(r[1])
        vago = (nome or "").upper() == "VAGO"
        proprietarios.append({"id": cod, "nome": nome or cod})
        imoveis.append({
            "id": f"{cod}1",
            "proprietario_id": cod,
            "status": "vago" if vago else "disponivel",
            "contato_telefone": s(r[2]),
        })
    return proprietarios, imoveis


def ler_comissoes(wb, contratos_validos=None, competencia_padrao="2026-09-01"):
    """
    Aba 'Pgto Comissões Gerais (2)'.
    D = Valor do Aluguel ; M = Valor Recorrente (= aluguel × 10%)
    G..K = comissões do 1º aluguel ; N..R = comissões do recorrente
    Recalcula a partir do RATEIO em vez de copiar as células, para que
    o banco fique consistente com a regra e não com arredondamentos soltos.
    """
    ws = wb[ABA_COMISSAO]
    comp = dt(ws["C2"].value) or competencia_padrao
    out = []

    # A aba tem TRÊS blocos empilhados que reaproveitam os mesmos códigos
    # na coluna A, com significados diferentes nas mesmas colunas:
    #    4..26  comissões (este bloco)
    #    29     linha de totais
    #    31..40 resumo por pessoa
    #    42..   'PAGAMENTOS PARA PROPRIETÁRIOS' (coluna E é IPTU, não captador)
    # Ler além da linha 26 gerava comissões fabricadas a partir de valores
    # de condomínio e IPTU.
    orfaos = {}
    for row in ws.iter_rows(min_row=4, max_row=26, values_only=True):
        cod = s(row[0])
        if not cod or not re.match(r"^\d{4}-\d{3}(\.\d+)?$", cod):
            continue
        if contratos_validos is not None and cod not in contratos_validos:
            orfaos[cod] = (sum(num(row[i]) for i in range(6, 11))
                           + num(row[12]))
            continue

        aluguel = num(row[3])
        # Sem fallback para aluguel × 10%: quando a coluna M está vazia, o
        # contrato está no primeiro mês e ainda não gera recorrente. Calcular
        # mesmo assim inventava comissão que a planilha não paga.
        base_recorrente = num(row[12])
        base_primeiro = sum(num(row[i]) for i in range(6, 11))  # G..K

        corretor, captador = s(row[4]), s(row[5])
        nomes = {
            "Corretor": corretor,
            "Captador": captador,
            "Gestor": "VALDIR",
            "Gestor ADM": None,
            "Imobiliaria": "BEEMOB",
        }

        for funcao, (pct_prim, pct_rec) in RATEIO.items():
            for tipo, base, pct in (
                ("primeiro", base_primeiro, pct_prim),
                ("recorrente", base_recorrente, pct_rec),
            ):
                valor = round(base * pct, 2)
                if valor <= 0:
                    continue
                out.append({
                    "contrato_id": cod,
                    "corretor_nome": nomes.get(funcao),
                    "funcao": funcao,
                    "tipo": tipo,
                    "base_calculo": round(base, 2),
                    "percentual": pct,
                    "valor": valor,
                    "competencia": comp,
                    "pago": False,
                })

    for cod, valor in sorted(orfaos.items()):
        print(f"   ⚠ código {cod} tem R$ {brl(valor)} de comissão na planilha mas "
              f"não existe na aba 'Cadastros' — NÃO importado. "
              f"O cliente precisa cadastrá-lo para o valor entrar no sistema.")
    return out


def conferir_totais_comissao(wb, comissoes):
    """
    O rodapé da aba de comissões traz os totais por pessoa que o cliente
    usa para pagar. Confere contra o que foi recalculado a partir do rateio,
    para que qualquer divergência apareça antes de gravar.
    """
    ws = wb[ABA_COMISSAO]

    # Linha 40 da aba: 'TOTAIS' do resumo por pessoa.
    # B = 1º aluguel, C = recorrente. É o número que o cliente usa para pagar.
    if s(ws["A40"].value) != "TOTAIS":
        print("   · conferência de comissões: linha de totais não encontrada, pulando")
        return

    alvo_primeiro = num(ws["B40"].value)
    alvo_recorrente = num(ws["C40"].value)

    calc_primeiro = sum(c["valor"] for c in comissoes if c["tipo"] == "primeiro")
    calc_recorrente = sum(c["valor"] for c in comissoes if c["tipo"] == "recorrente")

    for rotulo, calc, alvo in (
        ("1º aluguel", calc_primeiro, alvo_primeiro),
        ("recorrente", calc_recorrente, alvo_recorrente),
    ):
        dif = calc - alvo
        marca = "✓" if abs(dif) <= 0.5 else "⚠"
        print(f"   {marca} comissão {rotulo}: recalculado R$ {brl(calc)} · "
              f"planilha R$ {brl(alvo)}" +
              (f" · diferença R$ {brl(dif)}" if abs(dif) > 0.5 else ""))

    # A planilha paga 'ÉRIKA / JUR/ADM', que não existe no rateio de 5 funções
    for r in range(31, 40):
        nome = s(ws.cell(r, 1).value)
        total = num(ws.cell(r, 4).value)
        if nome and "JUR" in nome.upper() and total:
            print(f"     ⚠ '{nome}' recebe R$ {brl(total)} na planilha mas não faz "
                  f"parte do rateio de comissão — confirmar com o cliente o que é")


def ler_xlsm(path: Path):
    """
    Uma planilha financeira = um imóvel.
    'Cadastros Gerais' linha 3 = proprietário; linhas 8+ = imóveis/inquilinos.
    """
    wbf = openpyxl.load_workbook(path, data_only=False, keep_vba=True)
    wbv = openpyxl.load_workbook(path, data_only=True, keep_vba=True)
    cg = wbv["Cadastros Gerais"]

    cod = s(cg["A3"].value)
    if not cod:
        print(f"   ! {path.name}: sem código de proprietário em A3, pulando")
        return None

    proprietario = {
        "id": cod,
        "nome": s(cg["B3"].value),
        "cpf": s(cg["E3"].value),
        "rg": s(cg["F3"].value),
        "orgao_emissor": s(cg["G3"].value),
        "data_nasc": dt(cg["H3"].value),
        "naturalidade": s(cg["I3"].value),
        "estado_civil": s(cg["K3"].value),
        "telefone": s(cg["L3"].value),
        "endereco_proprietario": s(cg["M3"].value),
        "cep_proprietario": s(cg["R3"].value),
        "cidade_uf_proprietario": s(cg["S3"].value),
        "email": s(cg["T3"].value),
        "banco": s(cg["W3"].value),
        "agencia": s(cg["X3"].value),
        "conta_corrente": s(cg["Y3"].value),
        "pix": s(cg["Z3"].value),
    }

    imoveis, inquilinos, contratos = [], [], []
    for linha in range(8, 13):  # a planilha reserva 5 slots de imóvel
        imovel_id = s(cg[f"A{linha}"].value)
        aluguel = num(cg[f"G{linha}"].value)
        if not imovel_id or aluguel <= 0:
            continue

        imoveis.append({
            "id": imovel_id,
            "proprietario_id": cod,
            "inscricao_municipal": s(cg[f"B{linha}"].value),
            "endereco": s(cg[f"AS{linha}"].value),
            "cep": s(cg[f"AX{linha}"].value),
            "cidade_uf": s(cg[f"AY{linha}"].value),
            "valor_aluguel": aluguel,
            "valor_iptu": num(cg[f"H{linha}"].value),
            "dia_vencto_iptu": inteiro(cg[f"I{linha}"].value),
            "valor_condominio": num(cg[f"J{linha}"].value),
            "taxas_extras_condominio": num(cg[f"K{linha}"].value),
            "dia_vencto_condominio": inteiro(cg[f"L{linha}"].value),
            "administradora": s(cg[f"M{linha}"].value),
            "taxas_extras_manutencao": num(cg[f"P{linha}"].value),
            "garantia": s(cg[f"Q{linha}"].value),
            "valor_caucao": num(cg[f"R{linha}"].value),
            "percentual_adm_caucao": num(cg[f"S{linha}"].value),
            "valor_parcial_caucao": num(cg[f"T{linha}"].value),
            "valor_total_com_adm_caucao": num(cg[f"U{linha}"].value),
            "captador_nome": s(cg[f"V{linha}"].value),
            "corretor_nome": s(cg[f"W{linha}"].value),
            "gestor_nome": s(cg[f"X{linha}"].value),
            "imobiliaria_nome": s(cg[f"Y{linha}"].value) or "BEEMOB",
            "honorario_adm_primeiro": num(cg[f"Z{linha}"].value),
            "primeiro_pgto_hon": num(cg[f"AA{linha}"].value),
            "data_primeiro_pgto": dt(cg[f"AB{linha}"].value),
            "segundo_pgto_hon": num(cg[f"AC{linha}"].value),
            "data_segundo_pgto": dt(cg[f"AD{linha}"].value),
            "percentual_taxa_adm": num(cg[f"AF{linha}"].value, TAXA_ADM_PADRAO),
            "valor_adm": num(cg[f"AG{linha}"].value),
        })

        if s(cg[f"AH{linha}"].value):
            inquilinos.append({
                "id": imovel_id,
                "nome": s(cg[f"AH{linha}"].value),
                "cpf": s(cg[f"AK{linha}"].value),
                "rg": s(cg[f"AL{linha}"].value),
                "orgao_emissor": s(cg[f"AM{linha}"].value),
                "data_nasc": dt(cg[f"AN{linha}"].value),
                "naturalidade": s(cg[f"AO{linha}"].value),
                "estado_civil": s(cg[f"AQ{linha}"].value),
                "telefone": s(cg[f"AR{linha}"].value),
                "endereco": s(cg[f"AS{linha}"].value),
                "cep": s(cg[f"AX{linha}"].value),
                "cidade_uf": s(cg[f"AY{linha}"].value),
                "email": s(cg[f"AZ{linha}"].value),
            })

        contratos.append({
            "id": cod,
            "proprietario_id": cod,
            "inquilino_id": imovel_id,
            "imovel_id": imovel_id,
            "localizacao": s(cg[f"AS{linha}"].value),
            "data_inicio": dt(cg[f"C{linha}"].value),
            "data_fim": dt(cg[f"D{linha}"].value),
            "prazo_meses": inteiro(cg[f"E{linha}"].value) or 30,
            "dia_pagamento": inteiro(cg[f"F{linha}"].value) or 10,
            "garantia": s(cg[f"Q{linha}"].value),
            "valor_recorrente": aluguel,
            "percentual_taxa": num(cg[f"AF{linha}"].value, TAXA_ADM_PADRAO),
            "captador_nome": s(cg[f"V{linha}"].value),
            "corretor_nome": s(cg[f"W{linha}"].value),
            "gestor_nome": s(cg[f"X{linha}"].value),
            "status": "ativo",
        })

    inicio_contrato = contratos[0]["data_inicio"] if contratos else None
    fim_contrato = contratos[0]["data_fim"] if contratos else None
    lancamentos = ler_recibos(wbf, wbv, cod, imoveis, inicio_contrato, fim_contrato)
    itens_cond = ler_tx_condominio(wbv, imoveis)
    wbf.close()
    wbv.close()

    return {
        "proprietario": proprietario,
        "imoveis": imoveis,
        "inquilinos": inquilinos,
        "contratos": contratos,
        "lancamentos": lancamentos,
        "condominio_itens": itens_cond,
    }


def ler_recibos(wbf, wbv, cod, imoveis, inicio_contrato=None, fim_contrato=None):
    """
    Abas 'Rec_MMAAAA' = demonstrativo mensal.

    O layout varia: a planilha da Nádia tem o líquido em H19, a da Kamila
    em H20 (ela ganhou uma linha de SEGURO FIANÇA). Por isso as linhas são
    localizadas pelo rótulo da coluna A, não por posição fixa.
    """
    imovel_id = imoveis[0]["id"] if imoveis else f"{cod}1"
    out = []

    for nome in wbv.sheetnames:
        m = re.match(r"^Rec_(\d{2})(\d{4})$", nome)
        if not m:
            continue
        wsv = wbv[nome]

        rotulos = {}
        for r in range(9, 26):
            rot = s(wsv.cell(r, 1).value)
            if rot:
                rotulos[rot.lower()] = r

        def valor_de(*chaves, col=8):
            for chave in chaves:
                for rot, r in rotulos.items():
                    if chave in rot:
                        return num(wsv.cell(r, col).value)
            return 0.0

        def linha_de(*chaves):
            for chave in chaves:
                for rot, r in rotulos.items():
                    if chave in rot:
                        return r
            return None

        # Período de cálculo: F/G da linha do aluguel
        r_aluguel = linha_de("aluguel referente", "aluguél referente")
        if not r_aluguel:
            continue
        ini = wsv.cell(r_aluguel, 6).value
        fim = wsv.cell(r_aluguel, 7).value
        if not isinstance(ini, (date, datetime)):
            continue
        ini_d = ini.date() if isinstance(ini, datetime) else ini
        fim_d = (fim.date() if isinstance(fim, datetime) else fim) if isinstance(fim, (date, datetime)) \
            else date(ini_d.year, ini_d.month, dias_no_mes(ini_d))

        comp = wsv["C3"].value
        comp_d = comp.date() if isinstance(comp, datetime) else (comp if isinstance(comp, date) else None)

        aluguel = valor_de("aluguel referente", "aluguél referente")
        condominio = valor_de("taxa de condomínio")
        iptu = valor_de("iptu referente")
        seguro = valor_de("seguro fiança")
        extraordinaria = valor_de("extraordinária")
        honorarios = valor_de("honorários de administração")
        taxa_adm = valor_de("taxa administração", "taxa de administração imobiliária")
        manutencao = valor_de("despesas extras")
        juros = valor_de("juros e multa")

        # Complemento de condomínio (K12 na Kamila): fora do fluxo do inquilino
        r_cond = linha_de("taxa de condomínio")
        complemento = num(wsv.cell(r_cond, 11).value) if r_cond else 0.0

        subtotal = aluguel + condominio + iptu + seguro
        repasse = (subtotal - condominio - complemento - iptu - seguro
                   - extraordinaria - honorarios - taxa_adm - manutencao + juros)

        # Confere contra o líquido que a própria planilha calculou
        r_liq = linha_de("valor líquido para proprietário")
        liquido_planilha = num(wsv.cell(r_liq, 8).value) if r_liq else None
        ajuste = 0.0
        motivo = None
        if liquido_planilha is not None and abs(liquido_planilha - repasse) > 0.01:
            ajuste = round(liquido_planilha - repasse, 2)
            motivo = f"Diferença vs. fórmula da aba {nome} (valor cravado na planilha)"
            repasse = liquido_planilha

        out.append({
            "_aba": nome,
            "_competencia": comp_d.isoformat() if comp_d else None,
            "contrato_id": cod,
            "imovel_id": imovel_id,
            "periodo_inicio": ini_d.isoformat(),
            "periodo_fim": fim_d.isoformat(),
            "dias_periodo": (fim_d - ini_d).days + 1,
            "dias_mes": dias_no_mes(ini_d),
            "e_primeiro_mes": honorarios > 0,
            "valor_aluguel": round(aluguel, 2),
            "valor_condominio": round(condominio, 2),
            "complemento_condominio": round(complemento, 2),
            "valor_iptu": round(iptu, 2),
            "seguro_fianca": round(seguro, 2),
            "taxa_extraordinaria_cond": round(extraordinaria, 2),
            "honorarios_primeiro_aluguel": round(honorarios, 2),
            "taxa_adm_imobiliaria": round(taxa_adm, 2),
            "despesas_manutencao": round(manutencao, 2),
            "juros_multa": round(juros, 2),
            "subtotal_inquilino": round(subtotal, 2),
            "valor_liquido_inquilino": round(subtotal, 2),
            "valor_repasse_proprietario": round(repasse, 2),
            "ajuste_manual": ajuste,
            "ajuste_motivo": motivo,
            "status_pagamento": "pago",
        })

    return corrigir_periodos(out, inicio_contrato, fim_contrato)


def corrigir_periodos(lancs, inicio_contrato, fim_contrato):
    """
    Segunda passagem sobre os períodos.

    A planilha da Nádia tem a aba 'Rec_122026' com F11 = 2026-12-10, enquanto
    a competência (C3) diz 2025-12-01 e o contrato começa em 2025-12-10 — o
    ano foi digitado errado na origem. O efeito é que o primeiro mês do
    contrato some da série e aparece um mês solto lá na frente.

    Regra: corrige pela competência C3 só quando ela é coerente (cai dentro
    da vigência) e nenhum outro lançamento já cobre aquele mês. Isso evita
    corrigir as abas 08 e 09, cuja competência ficou desatualizada em
    2026-07 mas cujo período F11 está certo.
    """
    if not inicio_contrato:
        for l in lancs:
            l.pop("_aba", None)
            l.pop("_competencia", None)
        return lancs

    ini_ct = date.fromisoformat(inicio_contrato)
    fim_ct = date.fromisoformat(fim_contrato) if fim_contrato else None
    meses_ocupados = {l["periodo_inicio"][:7] for l in lancs}

    for l in lancs:
        aba = l.pop("_aba", None)
        comp = l.pop("_competencia", None)
        if not comp:
            continue
        comp_d = date.fromisoformat(comp)
        mes_periodo = l["periodo_inicio"][:7]
        mes_comp = comp[:7]

        if mes_comp == mes_periodo:
            continue
        if mes_comp in meses_ocupados:
            continue                       # outra aba já cobre esse mês
        if comp_d < date(ini_ct.year, ini_ct.month, 1):
            continue
        if fim_ct and comp_d > fim_ct:
            continue

        novo_ini = max(ini_ct, date(comp_d.year, comp_d.month, 1))
        novo_fim = date(comp_d.year, comp_d.month, dias_no_mes(comp_d))
        aviso = (f"{aba}: período {l['periodo_inicio']}..{l['periodo_fim']} "
                 f"inconsistente com a competência {mes_comp}; "
                 f"corrigido para {novo_ini}..{novo_fim}")
        meses_ocupados.discard(mes_periodo)
        meses_ocupados.add(mes_comp)
        l["periodo_inicio"] = novo_ini.isoformat()
        l["periodo_fim"] = novo_fim.isoformat()
        l["dias_periodo"] = (novo_fim - novo_ini).days + 1
        l["dias_mes"] = dias_no_mes(novo_ini)
        l["ajuste_motivo"] = f"{l['ajuste_motivo'] + ' | ' if l['ajuste_motivo'] else ''}{aviso}"

    return lancs


def ler_tx_condominio(wbv, imoveis):
    """Aba 'Tx Condom.': colunas D (descrição) e E (valor)."""
    if "Tx Condom." not in wbv.sheetnames or not imoveis:
        return []
    ws = wbv["Tx Condom."]
    imovel_id = imoveis[0]["id"]
    out = []
    for r in range(1, ws.max_row + 1):
        desc = s(ws.cell(r, 4).value)
        valor = num(ws.cell(r, 5).value, None)
        if not desc or valor is None:
            continue
        out.append({
            "imovel_id": imovel_id,
            "competencia": date.today().replace(day=1).isoformat(),
            "descricao": desc,
            "valor": round(valor, 2),
            "extraordinaria": "extra" in desc.lower(),
        })
    return out


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Importa as planilhas da Beemob para o Supabase")
    ap.add_argument("--commit", action="store_true",
                    help="grava no Supabase (sem esta flag, apenas simula)")
    args = ap.parse_args()

    url, key = env_supabase()
    if not url:
        sys.exit("✗ VITE_SUPABASE_URL não encontrada (.env ou variável de ambiente)")
    if args.commit and not key:
        sys.exit("✗ SUPABASE_SERVICE_KEY não definida no ambiente.\n"
                 "  Supabase > Project Settings > API > service_role\n"
                 '  PowerShell:  $env:SUPABASE_SERVICE_KEY = "..."')

    modo = "GRAVANDO" if args.commit else "DRY-RUN (nada será gravado)"
    print(f"\n{'='*64}\nBEEMOB — importação de planilhas · {modo}\n{'='*64}")

    supa = Supa(url, key, args.commit)

    # ---- 1. Consolidado -------------------------------------------------
    if not XLSX_COMISSAO.exists():
        sys.exit(f"✗ não encontrei {XLSX_COMISSAO}")

    print(f"\n[1/3] {XLSX_COMISSAO.name}")
    wb = openpyxl.load_workbook(XLSX_COMISSAO, data_only=True)

    corretores = ler_corretores(wb)
    print(f"   · {len(corretores)} corretores")
    supa.upsert("corretores", corretores)

    props, inqs, imoveis, contratos = ler_cadastros(wb)
    print(f"   · {len(props)} proprietários · {len(imoveis)} imóveis alugados · "
          f"{len(inqs)} inquilinos · {len(contratos)} contratos")

    props_disp, imoveis_disp = ler_imoveis_disponiveis(wb)
    vagos = sum(1 for i in imoveis_disp if i["status"] == "vago")
    print(f"   · {len(imoveis_disp)} imóveis na carteira sem contrato "
          f"({vagos} marcados VAGO)")

    # O mesmo código pode aparecer nos dois blocos com proprietários
    # diferentes. Nesse caso o bloco principal (com contrato) vence.
    por_id = {p["id"]: p["nome"] for p in props}
    conflitos = [(p["id"], por_id[p["id"]], p["nome"]) for p in props_disp
                 if p["id"] in por_id and por_id[p["id"]] != p["nome"]]
    for cid, n1, n2 in conflitos:
        print(f"   ⚠ CONFLITO de código {cid}: '{n1}' (com contrato) vs "
              f"'{n2}' (para alugar) — mantido o primeiro, confirmar com o cliente")
    ids_conflito = {c[0] for c in conflitos}
    props_disp = [p for p in props_disp if p["id"] not in ids_conflito]
    imoveis_disp = [i for i in imoveis_disp if i["proprietario_id"] not in ids_conflito]

    supa.upsert("proprietarios", props + props_disp)
    supa.upsert("inquilinos", inqs)
    supa.upsert("imoveis", imoveis + imoveis_disp)
    supa.upsert("contratos", contratos)

    caucoes = ler_caucoes(wb)
    total_caucao = sum(c["valor_total"] for c in caucoes)
    print(f"   · {len(caucoes)} cauções sob custódia — total R$ {brl(total_caucao)}")
    supa.upsert("caucoes", caucoes, on_conflict="contrato_id")

    comissoes = ler_comissoes(wb, contratos_validos={c["id"] for c in contratos})
    print(f"   · {len(comissoes)} linhas de comissão")
    conferir_totais_comissao(wb, comissoes)
    wb.close()

    # ---- 2. Planilhas financeiras detalhadas ----------------------------
    arquivos = sorted(DOWNLOADS.glob(XLSM_GLOB))
    print(f"\n[2/3] {len(arquivos)} planilha(s) financeira(s) detalhada(s)")

    for path in arquivos:
        print(f"\n   → {path.name}")
        dados = ler_xlsm(path)
        if not dados:
            continue
        print(f"     {len(dados['imoveis'])} imóvel(is) · "
              f"{len(dados['lancamentos'])} lançamentos · "
              f"{len(dados['condominio_itens'])} itens de condomínio")
        supa.upsert("proprietarios", [dados["proprietario"]])
        supa.upsert("imoveis", dados["imoveis"])
        supa.upsert("inquilinos", dados["inquilinos"])
        supa.upsert("contratos", dados["contratos"])
        supa.upsert("lancamentos", dados["lancamentos"],
                    on_conflict="contrato_id,periodo_inicio")
        supa.upsert("condominio_itens", dados["condominio_itens"],
                    on_conflict="imovel_id,competencia,descricao")

        for l in dados["lancamentos"]:
            if l["ajuste_motivo"]:
                valor = (f"ajuste de R$ {l['ajuste_manual']:.2f} — "
                         if l["ajuste_manual"] else "")
                print(f"     ⚠ {l['periodo_inicio']}: {valor}{l['ajuste_motivo']}")

    # ---- 3. Comissões ---------------------------------------------------
    print(f"\n[3/3] comissões")
    supa.upsert("comissoes", comissoes, on_conflict="contrato_id,funcao,tipo,competencia")

    print(f"\n{'='*64}")
    for tabela, n in sorted(supa.contagem.items()):
        print(f"  {tabela:20} {n:>5}")
    print("=" * 64)
    if not args.commit:
        print("Dry-run concluído. Rode com --commit para gravar.\n")
    else:
        print("Importação concluída.\n")


if __name__ == "__main__":
    main()
